"""Build a native, editable product review and an identical SVG web presentation."""

from html import escape
from pathlib import Path
import json

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
PUBLIC = ROOT / "public"
IMAGES = PUBLIC / "deck-images"
PPTX = PUBLIC / "Mathify-stakeholder-evidence.pptx"
INK = "#202124"
GRAY = "#5f6368"
BLUE = "#1a73e8"
GREEN = "#13795b"
AMBER = "#a14400"
LINE = "#dadce0"
PALE = "#f8fafd"
BLUE_PALE = "#eaf2ff"
GREEN_PALE = "#e7f4ed"
AMBER_PALE = "#fff1e5"
FONT = "Segoe UI"
FONT_PATH = Path(r"C:\Windows\Fonts")
W, H = 1600, 900
EEF = "https://educationendowmentfoundation.org.uk/education-evidence/guidance-reports/feedback"
WEI = "https://arxiv.org/abs/2201.11903"
GOOGLE = "https://services.google.com/fh/files/misc/goldman_sachs_keynote_google_cloud_thomas_kurian_021120.pdf"
SPARKS = "https://www.wearesparks.com/our-work/google-cloud-next/"
TEMPLATE = "https://www.slidescarnival.com/template/beatrice-free-presentation-template/66"
SLIDES = [
    {
        "section": "PRODUCT REVIEW",
        "title": "A clearer starting point\nfor tutoring.",
        "script": "The idea is to make the start of a tutoring session more useful. A student's question tells us what they're working on, but it doesn't always tell us where they're stuck. Mathify uses the waiting time to ask short concept questions, then puts those answers in a brief the tutor can inspect. Today I'll walk through one example, explain why the approach is worth testing, and show exactly which evidence we have and which evidence is still missing.",
    },
    {
        "section": "THE OPPORTUNITY",
        "title": "The question shows the task.\nIt does not show the gap.",
        "script": "Take this equation. The student might be stuck on brackets, balancing, or dividing at the end. Looking at the question alone doesn't distinguish those possibilities. A short probe can give the tutor something more specific to discuss. It still isn't a diagnosis: a wrong choice could be a slip or a guess. We're trying to surface a useful starting point that the tutor confirms with the student, while leaving the whiteboard available for their own work.",
    },
    {
        "section": "RESEARCH RATIONALE",
        "title": "Check understanding\nbefore choosing the feedback.",
        "source": "Design rationale: EEF, Teacher Feedback to Improve Pupil Learning, 2021.",
        "links": [["EEF feedback guidance", EEF]],
        "script": "This is the strongest educational reason for the design. The EEF's feedback guidance recommends using formative assessment to establish learning gaps before giving feedback. We've translated that into a simple sequence: ask about a concept, inspect the response, and let the tutor decide where to begin. The research supports that principle. It does not evaluate Mathify, and it doesn't establish that our AI questions are accurate or that this brief saves tutor time. Those are separate things we need to measure.",
    },
    {
        "section": "PRODUCT FLOW",
        "title": "One problem becomes\na traceable tutor handoff.",
        "script": "The student submits a problem, and the AI breaks it into ordered concepts. For each concept, the app asks a main probe and has an easier retry ready. A wrong answer or 'I don't know' leads to that retry. The brief keeps the actual responses, including early submission and concepts that weren't reached. The tutor can inspect the gap rather than trust an unexplained summary. Whiteboard work travels with the handoff as an image preview; the current AI does not assess that handwriting.",
    },
    {
        "section": "SYNTHETIC WALKTHROUGH",
        "title": "A missed bracket question\nleads to a simpler check.",
        "source": "Actual prototype screenshot with synthetic answers. Not a live-model benchmark.",
        "script": "Here's the same example all the way through. The main probe asks the student to expand 2 times x plus 3. They choose 2x plus 3, which misses multiplying the second term. The easier question asks about 2 times x plus 1, and they choose the correct 2x plus 2. That suggests a simpler example may be a useful starting point, but the tutor should check their thinking. This is a synthetic walkthrough, not an actual student or live-model accuracy result.",
    },
    {
        "section": "TUTOR EXPERIENCE",
        "title": "Put the useful detail first.",
        "source": "Actual priority card from the same synthetic walkthrough.",
        "script": "This is the priority card in the working tutor view. The possible gap comes first. The tutor can compare the selected answer with the key, and see that the student got the easier retry right. That supports a specific conversation about how multiplication applies to both terms. It doesn't mean the student has a confirmed misconception. The confidence label elsewhere in the brief is also a heuristic, not calibrated accuracy. Our next usability check is whether tutors can read this quickly and interpret it correctly.",
    },
    {
        "section": "QUALITY SAFEGUARDS",
        "title": "A second AI check helps.\nIt is not ground truth.",
        "links": [["Wei et al., 2022: language-model reasoning context", WEI]],
        "script": "There are different kinds of checks, and they answer different questions. Schema validation checks that the response has a usable structure. The optional second AI call reviews the breakdown and the probes, including the easier question's hint and explanation. But two calls to the same provider can agree on a mistake. An independent mathematical and educator review is still needed. We have reference answers and a harness to start that work, but we haven't run a live accuracy benchmark. No live-model accuracy percentage is claimed.",
    },
    {
        "section": "CURRENT EVIDENCE",
        "title": "The workflow is tested.\nEffectiveness is still unmeasured.",
        "source": "Engineering snapshot from this session. Mock fixtures do not measure model accuracy.",
        "script": "Here's the boundary of what we've verified. The offline suite has 111 passing tests, and three browser checks cover the synthetic handoff flows and the presentation. Eight reference solutions have also been checked. These numbers tell us about software behavior and reference mathematics, not the quality of 111 generated questions or outcomes for 111 students. Live probe correctness, tutor comprehension, and student benefit remain unmeasured.\n\nChecked reference answers: x = 4; (x + 2)(x + 3); hypotenuse 5 cm; sin θ = 5/13; 5 cups; (x, y) = (2, 5); circle area 153.9 cm²; x⁴; x ≠ 0. These are reference answers, not generated-model results.",
    },
    {
        "section": "TUTOR FEEDBACK",
        "title": "The questionnaire is out.\nThe responses are still pending.",
        "source": "Zero tutor responses so far. In-app form shown with synthetic session data.",
        "script": "There are zero tutor responses so far, so there's no tutor feedback analysis to present yet. The external questionnaire has been distributed and is intended to tell us about tutors' needs and layout preferences. This in-app form asks about an individual report: usefulness, breakdown accuracy, and probe quality. We checked the submission path and fixed the false-success behavior on failed saves. Once responses arrive, we'll report the count, rating distributions, and comments without overgeneralizing a small sample. A timed walkthrough will complement those opinions with observed behavior.",
    },
    {
        "section": "VALIDATION PLAN",
        "title": "Validate correctness,\nthen test usefulness.",
        "source": "Proposed evaluation measures, not results. Agree success thresholds before testing.",
        "script": "I'd validate this in stages. First, generate live items for the existing reference cases and have a reviewer check the concepts, answer keys, distractors, hints, and retries. Record the number reviewed and the errors, rather than just an overall success claim. Then give tutors matched synthetic cases with and without the brief, counterbalancing which one they see first. Measure time to the main takeaway and whether that takeaway is correct, alongside their feedback. A small student pilot should follow only if the mathematical and tutor-review gates are met.",
    },
    {
        "section": "NEXT MILESTONE",
        "title": "A tutor-reviewed demo,\nnot another feature sprint.",
        "script": "The next milestone I'd ask us to approve is a small validation round. We need live-model access, someone to independently review the mathematics, and tutors for a short walkthrough. Alec already put the breakdown, reviewer, evaluation harness, and survey scaffolding in place. The work in this session extends the whiteboard, makes feedback saving reliable, and makes the evidence and handoff easier to inspect. The remaining question isn't whether we can build more screens. It's whether tutors get a correct, useful starting point from this one.",
    },
    {
        "section": "APPENDIX",
        "title": "Sources, scope, and design references.",
        "links": [
            ["EEF: formative assessment and feedback", EEF],
            ["Wei et al., 2022: model reasoning, not student outcomes", WEI],
            ["Ontario MTH1W: curriculum context", "https://www.dcp.edu.gov.on.ca/en/curriculum/secondary-mathematics/courses/mth1w"],
            ["Public Google Cloud executive deck, February 2020", GOOGLE],
            ["Sparks: Google Cloud Next keynote design case study", SPARKS],
            ["Beatrice template lineage, SlidesCarnival", TEMPLATE],
            ["CC BY 4.0", "https://creativecommons.org/licenses/by/4.0/"],
        ],
        "script": "These references serve different purposes. EEF supports the educational rationale. Wei and colleagues provide context for language-model reasoning research, not evidence of student learning in Mathify. Ontario provides curriculum context, not verified textbook attribution. For the presentation design, I inspected an actual public Google Cloud executive deck and the Google Cloud Next case study from Google's design partner, Sparks. The layouts here are original, inspired by the hierarchy, spacing, and diagram-led storytelling, with no Google branding or endorsement. The PowerPoint retains its Beatrice template lineage and attribution under CC BY 4.0. The screenshots use mock data.\n\nOriginal template resource credits: SlidesCarnival, Unsplash, and SubtlePatterns. The original photos and textured backgrounds are not used in the visible compositions.",
    },
]


class Canvas:
    def __init__(self):
        self.items = []

    def text(self, x, y, text, size=28, color=INK, bold=False, width=1456):
        self.items.append(dict(kind="text", x=x, y=y, text=text, size=size, color=color, bold=bold, width=width))

    def rect(self, x, y, width, height, fill="white", stroke=None, radius=0):
        self.items.append(dict(kind="rect", x=x, y=y, width=width, height=height, fill=fill, stroke=stroke, radius=radius))

    def line(self, x, y, x2, y2, color=LINE, weight=2):
        self.items.append(dict(kind="line", x=x, y=y, x2=x2, y2=y2, color=color, weight=weight))

    def circle(self, x, y, diameter, fill):
        self.items.append(dict(kind="circle", x=x, y=y, width=diameter, height=diameter, fill=fill))

    def image(self, x, y, width, height, name, alt):
        path = IMAGES / f"{name}.png"
        with Image.open(path) as image:
            scale = min(width / image.width, height / image.height)
            actual_w, actual_h = image.width * scale, image.height * scale
        self.items.append(dict(kind="image", x=x, y=y, width=actual_w, height=actual_h, name=name, alt=alt))


def header(c, data, number, background="white"):
    c.rect(0, 0, W, H, background)
    c.text(72, 40, f"MATHIFY   /   {data['section']}", 16, GRAY, width=1200)
    c.text(72, 100, data["title"], 50, width=1456)
    c.line(72, 829, 1528, 829)
    c.text(72, 847, data.get("source", "CSC308  /  Stakeholder review  /  November 2026"), 15, GRAY, width=1330)
    c.text(1470, 847, f"{number:02d}", 15, GRAY, width=58)


def tag(c, x, y, text, color=BLUE, fill=BLUE_PALE, width=210):
    c.rect(x, y, width, 38, fill, radius=10)
    c.text(x + 14, y + 6, text, 18, color, width=width - 28)


def make_scenes():
    scenes = []
    for i, data in enumerate(SLIDES, 1):
        c = Canvas()
        header(c, data, i, PALE if i in {1, 11} else "white")
        if i == 1:
            c.text(76, 310, "Turn a student's question into\na focused tutor handoff.", 32, GRAY, width=710)
            c.text(76, 684, "Research rationale\nWorking prototype\nValidation plan", 23, GRAY, width=650)
            c.rect(864, 266, 650, 196, "white", LINE, 16)
            tag(c, 892, 286, "STUDENT QUESTION")
            c.text(902, 348, "2(x + 3) = 14", 52, width=550)
            c.line(1189, 462, 1189, 512, BLUE, 3)
            c.rect(908, 512, 606, 253, BLUE_PALE, radius=16)
            c.text(940, 541, "TUTOR STARTING POINT", 18, BLUE, width=540)
            c.text(940, 590, "Check distribution\nbefore solving.", 34, width=530)
            c.text(940, 701, "A possible gap to confirm.", 22, GRAY, width=530)
        elif i == 2:
            c.rect(72, 310, 615, 422, PALE, radius=16)
            c.text(104, 341, "WHAT WE KNOW", 18, GRAY, width=550)
            c.text(104, 423, "2(x + 3) = 14", 56, width=550)
            c.text(104, 580, "The problem.\nNot yet the student's thinking.", 27, GRAY, width=550)
            c.text(796, 311, "WHAT WE NEED TO CHECK", 18, BLUE, width=730)
            for y, n, name, equation in [(384, "01", "Distribute", "2x + 6 = 14"), (509, "02", "Balance", "2x = 8"), (634, "03", "Divide", "x = 4")]:
                c.text(796, y, n, 22, BLUE, width=70)
                c.text(884, y, name, 29, width=295)
                c.text(1220, y, equation, 29, GRAY, width=305)
                c.line(796, y + 78, 1528, y + 78)
        elif i == 3:
            c.rect(72, 318, 690, 419, BLUE_PALE, radius=16)
            c.text(106, 354, "THE DESIGN FOUNDATION", 18, BLUE, width=610)
            c.text(106, 425, "Formative assessment\nhelps establish the gap\nbefore feedback.", 36, width=610)
            c.text(106, 653, "EEF guidance, 2021  /  Recommendation 1", 20, GRAY, width=610)
            for y, n, title, detail in [
                (328, "1", "Ask about one concept", "A short probe, not the full solution."),
                (464, "2", "Inspect the answer", "Keep the response, not just a score."),
                (600, "3", "Let the tutor confirm", "Choose the next explanation together."),
            ]:
                c.circle(834, y + 4, 44, BLUE_PALE)
                c.text(848, y + 6, n, 24, BLUE, width=30)
                c.text(902, y, title, 28, width=626)
                c.text(902, y + 48, detail, 22, GRAY, width=626)
            c.text(72, 772, "Supports the rationale. Does not prove this prototype works.", 22, GRAY)
        elif i == 4:
            labels = [("Question", "Text or image"), ("Breakdown", "Ordered concepts"), ("Probes", "Main + easier retry"), ("Responses", "Answers + work"), ("Tutor brief", "Priority + evidence")]
            for j, (name, detail) in enumerate(labels):
                x = 72 + j * 301
                c.rect(x, 329, 250, 158, BLUE_PALE if j in {1, 2} else PALE, radius=12)
                c.text(x + 22, 349, f"0{j + 1}", 18, BLUE, width=200)
                c.text(x + 22, 390, name, 28, width=222)
                c.text(x + 22, 443, detail, 19, GRAY, width=225)
                if j < 4:
                    c.text(x + 258, 389, "→", 28, BLUE, width=45)
            c.text(72, 565, "THE ADAPTIVE STEP", 18, GRAY)
            c.text(72, 614, "Main probe", 31, width=350)
            c.text(386, 614, "Wrong / IDK", 25, AMBER, width=250)
            c.text(645, 614, "→", 32, BLUE, width=60)
            c.text(733, 614, "One easier retry", 31, width=380)
            c.text(1178, 614, "→", 32, BLUE, width=60)
            c.text(1260, 614, "Both retained", 29, width=270)
            c.line(72, 704, 1528, 704)
            c.text(72, 736, "Metadata makes the path inspectable. Whiteboard work stays available.", 24, GRAY)
        elif i == 5:
            c.rect(72, 305, 638, 428, PALE, radius=16)
            tag(c, 103, 337, "MAIN PROBE")
            c.text(103, 404, "Expand 2(x + 3)", 39, width=575)
            c.text(103, 495, "Selected", 19, GRAY, width=240)
            c.text(407, 495, "Answer key", 19, GRAY, width=260)
            c.text(103, 540, "2x + 3", 48, AMBER, width=250)
            c.text(407, 540, "2x + 6", 48, GREEN, width=260)
            c.line(103, 635, 678, 635)
            c.text(103, 663, "Second term not multiplied.", 24, AMBER, width=575)
            c.text(733, 469, "→", 36, BLUE, width=70)
            tag(c, 844, 282, "EASIER RETRY")
            c.image(844, 338, 670, 410, "retry", "Actual easier probe: expand 2(x+1), selected 2x+2.")
            c.text(72, 769, "A correct retry is evidence of a useful starting point, not mastery.", 22, GRAY)
        elif i == 6:
            c.text(72, 298, "ACTUAL TUTOR VIEW", 18, BLUE)
            c.image(72, 327, 1456, 310, "priority", "Actual priority card: main miss, selected answer, key, and correct easier retry.")
            c.text(1275, 430, "Main miss.\nCorrect retry.", 27, GRAY, width=250)
            for x, n, title, detail in [
                (72, "01", "Possible gap", "Distribution, not a generic score."),
                (570, "02", "Visible evidence", "The selected answer beside the key."),
                (1070, "03", "Retry context", "A simpler example was answered."),
            ]:
                c.text(x, 662, n, 18, BLUE, width=420)
                c.text(x, 697, title, 28, width=440)
                c.text(x, 746, detail, 20, GRAY, width=456)
        elif i == 7:
            rows = [
                (318, "01", "Output structure", "Usable fields and options", "IMPLEMENTED", GREEN, GREEN_PALE),
                (453, "02", "Second AI review", "Checks breakdown and probes", "OPTIONAL", BLUE, BLUE_PALE),
                (588, "03", "Independent review", "Math correctness and teaching value", "PENDING", AMBER, AMBER_PALE),
            ]
            for y, n, title, detail, status, color, fill in rows:
                c.text(72, y + 16, n, 26, color, width=80)
                c.text(176, y, title, 32, width=520)
                c.text(176, y + 52, detail, 23, GRAY, width=870)
                tag(c, 1238, y + 15, status, color, fill, 290)
                c.line(72, y + 105, 1528, y + 105)
            c.text(72, 755, "Agreement between AI calls is not independently verified correctness.", 24, GRAY)
        elif i == 8:
            c.text(72, 302, "VERIFIED ENGINEERING", 19, GREEN)
            for x, number, label in [(72, "111", "Offline tests"), (351, "3", "Browser checks"), (634, "8", "Reference solutions")]:
                c.text(x, 366, number, 80, width=250)
                c.text(x, 480, label, 23, GRAY, width=250)
            c.text(72, 614, "Build and type checks pass.", 27, width=720)
            c.text(72, 676, "These validate behavior,\nnot learning outcomes.", 26, GRAY, width=720)
            c.line(940, 312, 940, 748)
            c.text(1013, 302, "NOT YET MEASURED", 19, AMBER, width=515)
            for y, name in [(390, "Live probe correctness"), (500, "Tutor comprehension"), (610, "Student benefit")]:
                c.circle(1015, y + 13, 10, AMBER)
                c.text(1048, y, name, 29, width=470)
                c.line(1013, y + 76, 1528, y + 76)
        elif i == 9:
            c.text(72, 323, "0", 154, BLUE, width=660)
            c.text(80, 506, "Tutor responses", 36, width=650)
            c.text(80, 591, "Questionnaire distributed.\nNo findings to report yet.", 28, GRAY, width=670)
            c.text(80, 730, "Collection first. Analysis second.", 23, GRAY, width=670)
            c.image(801, 290, 727, 492, "feedback", "Actual tutor feedback form: usefulness, breakdown accuracy, and probe quality.")
        elif i == 10:
            c.line(100, 356, 1492, 356, BLUE, 3)
            for x, number, title, subtitle, lines in [
                (72, "01", "Check the maths", "Live probes + independent review", ["Concept and answer-key validity", "Distractors, hints, and retries", "Error counts with denominators"]),
                (575, "02", "Test the brief", "Matched cases, with vs. without", ["Time to the main takeaway", "Correct interpretation of the gap", "Tutor ratings and comments"]),
                (1078, "03", "Pilot carefully", "Only after the review gates", ["Agree the success thresholds", "Observe actual use with tutors", "Revise from evidence"]),
            ]:
                c.circle(x, 332, 48, BLUE)
                c.text(x + 10, 340, number, 21, "white", width=39)
                c.text(x, 420, title, 33, width=450)
                c.text(x, 479, subtitle, 21, GRAY, width=450)
                for j, line in enumerate(lines):
                    c.text(x, 568 + j * 60, line, 21, width=450)
        elif i == 11:
            c.text(72, 310, "Approve a small validation round.", 33, BLUE, width=880)
            for y, n, title in [(419, "01", "Enable live-model evaluation"), (516, "02", "Assign an independent math reviewer"), (613, "03", "Recruit tutors for a short walkthrough")]:
                c.text(72, y, n, 23, BLUE, width=80)
                c.text(160, y, title, 28, width=780)
                c.line(72, y + 67, 922, y + 67)
            c.rect(1002, 310, 526, 431, "white", LINE, 16)
            c.text(1036, 348, "READINESS", 18, GRAY, width=450)
            c.text(1036, 423, "Build ready.", 41, GREEN, width=450)
            c.text(1036, 491, "Evidence pending.", 37, AMBER, width=450)
            c.text(1036, 629, "Keep sessions and profiles\nout of this validation round.", 24, GRAY, width=450)
        elif i == 12:
            c.text(72, 270, "PRODUCT EVIDENCE", 18, BLUE, width=690)
            for y, title, detail in [
                (327, "EEF, 2021", "Formative assessment and feedback rationale."),
                (443, "Wei et al., 2022", "Model reasoning research, not learner outcomes."),
                (559, "Ontario MTH1W", "Curriculum context, not textbook verification."),
                (675, "Local verification", "Mock tests, reference maths, and browser flows."),
            ]:
                c.text(72, y, title, 27, width=690)
                c.text(72, y + 43, detail, 21, GRAY, width=690)
            c.line(811, 270, 811, 786)
            c.text(870, 270, "DESIGN REFERENCES + ATTRIBUTION", 18, BLUE, width=660)
            for y, title, detail in [
                (327, "Google Cloud public deck, 2020", "Hierarchy, whitespace, and diagram-led layouts."),
                (443, "Sparks / Google Cloud Next", "A coherent visual system for complex ideas."),
                (559, "Original Mathify compositions", "Native diagrams and synthetic UI screenshots."),
                (675, "Beatrice / SlidesCarnival / CC BY 4.0", "Template lineage retained. No Google endorsement."),
            ]:
                c.text(870, y, title, 25, width=658)
                c.text(870, y + 43, detail, 20, GRAY, width=658)
        scenes.append(c)
    return scenes


def text_height(item):
    return len(item["text"].split("\n")) * item["size"] * 1.38 + 4


def validate_scene(canvas, index):
    for item in canvas.items:
        kind = item["kind"]
        if kind == "text":
            font = ImageFont.truetype(str(FONT_PATH / ("segoeuib.ttf" if item["bold"] else "segoeui.ttf")), item["size"])
            for line in item["text"].split("\n"):
                measured = font.getlength(line)
                if measured > item["width"] + 1:
                    raise ValueError(f"Slide {index}: text exceeds authored width: {line} ({measured:.0f} > {item['width']})")
            height = text_height(item)
        elif kind == "line":
            assert 0 <= item["x2"] <= W and 0 <= item["y2"] <= H
            continue
        else:
            height = item["height"]
        if item["x"] < 0 or item["y"] < 0 or item["x"] + item["width"] > W + 1 or item["y"] + height > H + 1:
            raise ValueError(f"Slide {index}: element outside slide: {item}")


def rgb(value):
    return RGBColor.from_string("FFFFFF" if value == "white" else value.lstrip("#"))


def inch(px):
    return Inches(px / 120)


def render_ppt(presentation, scenes):
    for slide_id in list(presentation.slides._sldIdLst):
        presentation.part.drop_rel(slide_id.rId)
        presentation.slides._sldIdLst.remove(slide_id)
    presentation.slide_width = inch(W)
    presentation.slide_height = inch(H)
    for index, (canvas, data) in enumerate(zip(scenes, SLIDES), 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[8])
        slide._element.set("showMasterSp", "0")
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        for item in canvas.items:
            kind, x, y = item["kind"], inch(item["x"]), inch(item["y"])
            if kind == "image":
                picture = slide.shapes.add_picture(str(IMAGES / f"{item['name']}.png"), x, y, inch(item["width"]), inch(item["height"]))
                picture._element.xpath(".//p:cNvPr")[0].set("descr", item["alt"])
            elif kind == "text":
                shape = slide.shapes.add_textbox(x, y, inch(item["width"]), inch(text_height(item)))
                frame = shape.text_frame
                frame.word_wrap = False
                frame.margin_top = frame.margin_bottom = frame.margin_left = frame.margin_right = 0
                for j, line in enumerate(item["text"].split("\n")):
                    p = frame.paragraphs[0] if j == 0 else frame.add_paragraph()
                    p.text = line
                    p.alignment = PP_ALIGN.LEFT
                    p.space_before = p.space_after = Pt(0)
                    p.line_spacing = Pt(item["size"] * .6 * 1.38)
                    for run in p.runs:
                        run.font.name = FONT
                        run.font.size = Pt(item["size"] * .6)
                        run.font.bold = item["bold"]
                        run.font.color.rgb = rgb(item["color"])
            elif kind == "line":
                shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y, inch(item["x2"]), inch(item["y2"]))
                shape.line.color.rgb = rgb(item["color"])
                shape.line.width = Pt(item["weight"] * .6)
            else:
                shape_type = MSO_SHAPE.OVAL if kind == "circle" else MSO_SHAPE.ROUNDED_RECTANGLE if item.get("radius") else MSO_SHAPE.RECTANGLE
                shape = slide.shapes.add_shape(shape_type, x, y, inch(item["width"]), inch(item["height"]))
                if kind == "rect" and item.get("radius"):
                    shape.adjustments[0] = min(.5, item["radius"] / min(item["width"], item["height"]))
                shape.fill.solid()
                shape.fill.fore_color.rgb = rgb(item["fill"])
                if item.get("stroke"):
                    shape.line.color.rgb = rgb(item["stroke"])
                    shape.line.width = Pt(1)
                else:
                    shape.line.fill.background()
        for shape in slide.shapes:
            for style in shape._element.xpath("./p:style"):
                shape._element.remove(style)
        notes = data["script"]
        if data.get("links"):
            notes += "\n\nReferences:\n" + "\n".join(f"{label}: {url}" for label, url in data["links"])
        slide.notes_slide.notes_text_frame.text = notes
    presentation.core_properties.title = "Mathify: stakeholder product review"
    presentation.core_properties.subject = "Research rationale, traceable handoff, and validation plan. Zero tutor responses."
    presentation.core_properties.comments = "Original layouts. Beatrice template lineage by SlidesCarnival, CC BY 4.0. Google design references only; no endorsement."
    presentation.save(PPTX)


def render_svg(canvas, index):
    content = []
    for item in canvas.items:
        kind, x, y = item["kind"], item["x"], item["y"]
        if kind == "text":
            lines = item["text"].split("\n")
            tspans = "".join(f'<tspan x="{x}" dy="{0 if j == 0 else item["size"] * 1.38}">{escape(line)}</tspan>' for j, line in enumerate(lines))
            content.append(f'<text x="{x}" y="{y + item["size"]}" font-size="{item["size"]}" font-weight="{600 if item["bold"] else 400}" fill="{item["color"]}">{tspans}</text>')
        elif kind == "line":
            content.append(f'<line x1="{x}" y1="{y}" x2="{item["x2"]}" y2="{item["y2"]}" stroke="{item["color"]}" stroke-width="{item["weight"]}"/>')
        elif kind == "circle":
            content.append(f'<circle cx="{x + item["width"] / 2}" cy="{y + item["height"] / 2}" r="{item["width"] / 2}" fill="{item["fill"]}"/>')
        elif kind == "image":
            content.append(f'<image x="{x}" y="{y}" width="{item["width"]}" height="{item["height"]}" href="deck-images/{item["name"]}.png"><title>{escape(item["alt"])}</title></image>')
        else:
            content.append(f'<rect x="{x}" y="{y}" width="{item["width"]}" height="{item["height"]}" rx="{item["radius"]}" fill="{item["fill"]}" stroke="{item["stroke"] or "none"}"/>')
    return f'<svg class="slide-canvas" viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg" role="img" aria-labelledby="title-{index:02d}" style="font-family:\'Segoe UI\',Arial,sans-serif">{"".join(content)}</svg>'


def render_html(scenes):
    sections = []
    for index, (canvas, data) in enumerate(zip(scenes, SLIDES), 1):
        refs = "".join(f'<li><a href="{escape(url)}">{escape(label)}</a></li>' for label, url in data.get("links", []))
        sections.append(f"""<section class="slide" id="slide-{index:02d}" aria-labelledby="title-{index:02d}">
<h2 id="title-{index:02d}" class="sr-only" tabindex="-1">{escape(data["title"].replace(chr(10), " "))}</h2>
{render_svg(canvas, index)}
<aside class="notes"><h3>Speaker script</h3>{''.join(f'<p>{escape(p)}</p>' for p in data["script"].split(chr(10) + chr(10)))}{f'<ul>{refs}</ul>' if refs else ''}</aside>
</section>""")
    html = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Mathify: stakeholder product review</title><style>
*{box-sizing:border-box}body{margin:0;background:#e8eaed;color:#202124;font-family:"Segoe UI",Arial,sans-serif}
.toolbar{min-height:60px;display:flex;justify-content:center;gap:12px;align-items:center;flex-wrap:wrap;padding:10px 16px}
button{font:inherit;font-size:14px;padding:8px 13px;border:1px solid #dadce0;border-radius:5px;background:white;color:#202124;cursor:pointer}
button:disabled{opacity:.4;cursor:default}a{color:#1764c0;text-underline-offset:3px}
button:focus-visible,a:focus-visible{outline:3px solid #1a73e8;outline-offset:3px}
main{width:min(1600px,calc(100vw - 32px),calc((100vh - 86px) * 1.777778));margin:0 auto 16px}
.slide{margin-bottom:24px;background:white}.js .slide[hidden]{display:none}
.slide-canvas{display:block;width:100%;height:auto;aspect-ratio:16/9;background:white}
.sr-only{position:absolute;width:1px;height:1px;padding:0;margin:-1px;overflow:hidden;clip:rect(0,0,0,0);white-space:nowrap;border:0}
.notes{display:none;padding:28px 40px;border-top:1px solid #dadce0;max-width:1100px}
.show-notes .notes{display:block}.notes h3{font-size:21px;margin:0 0 16px}.notes p{font-size:19px;line-height:1.65}
.notes ul{font-size:16px;line-height:1.8;overflow-wrap:anywhere}#counter{font-size:14px;min-width:108px;text-align:center}
.toolbar a{font-size:14px}@media(max-width:600px){main{width:100%;margin-bottom:0}.toolbar{gap:6px;padding:8px}button,.toolbar a{font-size:12px}.notes{padding:20px}.notes p{font-size:17px}}
@media print{@page{size:A4 landscape;margin:0}body{background:white}.toolbar{display:none}main{width:100%;margin:0}.slide,.js .slide[hidden]{display:block!important;width:297mm;height:210mm;margin:0;break-after:page;page-break-after:always}.slide:last-child{break-after:auto;page-break-after:auto}.slide-canvas{width:297mm;height:167.0625mm;max-height:none;margin-top:21.46875mm}.notes{display:none!important}}
</style></head><body><nav class="toolbar" aria-label="Presentation controls">
<button id="prev" type="button">Previous</button><span id="counter" aria-live="polite"></span><button id="next" type="button">Next</button>
<button id="notes" type="button" aria-pressed="false">Speaker notes</button><button id="print" type="button">Print / Save PDF</button>
<a href="Mathify-stakeholder-evidence.pptx" download>Download PowerPoint</a></nav><main aria-label="Stakeholder presentation">
""" + "".join(sections) + """
</main><script>
(() => {
  const slides = [...document.querySelectorAll(".slide")];
  const previous = document.getElementById("prev"), next = document.getElementById("next");
  const notes = document.getElementById("notes"), counter = document.getElementById("counter");
  let index = 0;
  document.documentElement.classList.add("js");
  function show(target, focus = false) {
    index = Math.max(0, Math.min(slides.length - 1, target));
    slides.forEach((slide, i) => { slide.hidden = i !== index; });
    previous.disabled = index === 0; next.disabled = index === slides.length - 1;
    counter.textContent = `Slide ${index + 1} of ${slides.length}`;
    history.replaceState(null, "", `#${slides[index].id}`);
    if (focus) slides[index].querySelector("h2").focus({ preventScroll: true });
    window.scrollTo(0, 0);
  }
  function fromHash() {
    const target = slides.findIndex(slide => `#${slide.id}` === location.hash);
    show(target < 0 ? 0 : target);
  }
  function toggleNotes() { notes.setAttribute("aria-pressed", String(document.body.classList.toggle("show-notes"))); }
  previous.addEventListener("click", () => show(index - 1, true));
  next.addEventListener("click", () => show(index + 1, true));
  notes.addEventListener("click", toggleNotes);
  document.getElementById("print").addEventListener("click", () => window.print());
  window.addEventListener("hashchange", fromHash);
  document.addEventListener("keydown", event => {
    if (event.altKey || event.ctrlKey || event.metaKey || event.target.closest("input,textarea,select,button,a")) return;
    if (["ArrowRight","PageDown","ArrowLeft","PageUp","Home","End","n","N","p","P"].includes(event.key)) event.preventDefault();
    if (event.key === "ArrowRight" || event.key === "PageDown") show(index + 1, true);
    if (event.key === "ArrowLeft" || event.key === "PageUp") show(index - 1, true);
    if (event.key === "Home") show(0, true);
    if (event.key === "End") show(slides.length - 1, true);
    if (event.key.toLowerCase() === "n") toggleNotes();
    if (event.key.toLowerCase() === "p") window.print();
  });
  fromHash();
})();
</script></body></html>"""
    (PUBLIC / "stakeholder-deck.html").write_text(html, encoding="utf-8")


def build():
    if len(SLIDES) != 12:
        raise ValueError("Expected twelve authored slides.")
    scenes = make_scenes()
    for index, canvas in enumerate(scenes, 1):
        validate_scene(canvas, index)
    if "\u2014" in json.dumps(SLIDES, ensure_ascii=False):
        raise ValueError("Speaker script contains an em dash.")
    render_ppt(Presentation(PPTX), scenes)
    render_html(scenes)
    print("Built 12 original 16:9 product-review layouts with editable shapes and speaker scripts.")


if __name__ == "__main__":
    build()
